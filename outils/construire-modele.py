"""Construit module/modele-parc-vehicules.xlsx (modèle d'import du parc).
Les colonnes équipements se remplissent selon la solution (tableau « Équipements par solution » de l'onglet Listes).
Usage : python3 outils/construire-modele.py   (puis recalculer avec LibreOffice pour mettre les valeurs en cache)"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

SORTIE = pathlib.Path(__file__).resolve().parent.parent / 'module' / 'modele-parc-vehicules.xlsx'
NAVY, GREEN, GREY, LINE = '1C3465', '7FA127', 'F2F4F8', 'D8DEE8'
F = lambda **k: Font(name='Arial', **k)
thin = Side(style='thin', color=LINE); box = Border(left=thin, right=thin, top=thin, bottom=thin)
N = 500  # lignes préparées
SOLUTIONS = ["CITIPAP", "CITI'ID", "CITI'Mission", "CITIPAV"]
TYPES = ["BOM 12 T", "BOM 16 T", "BOM 19 T", "BOM 26 T", "BOM 32 T", "BOM 7,5 T", "Bennette 3,5 T", "BOM bicomp 50/50", "BOM bicomp 70/30",
         "BOM préhension latérale", "BOM grue 19 T", "BOM grue 26 T", "BOM grue 32 T"]
EQUIP = ["Tablette", "Support de charge", "Calculateur K3", "Antenne cabine", "Antenne arrière (dent)", "Décodeur", "Boîte à boutons (BBV)", "NBAB"]
CABINE = 4  # 4 premiers = cabine (toutes solutions) ; 4 suivants = arrière (CITI'ID)
# Onglet Parc : A Immatriculation, B Agence, C Solution, D Type, E Marque/modèle, F..M équipements, N Commentaire
C_EQ = 6  # 1re colonne équipement (F), identique dans l'onglet Listes pour garder les mêmes lettres

wb = Workbook()
# ---------- Mode d'emploi ----------
me = wb.active; me.title = "Mode d'emploi"; me.sheet_view.showGridLines = False
me.column_dimensions['A'].width = 3; me.column_dimensions['B'].width = 112
H = lambda t: (t, F(bold=True, size=11, color=NAVY)); P = lambda t: (t, F(size=10))
lignes = [("Modèle de parc véhicules - module Matériel embarqué", F(bold=True, size=14, color=NAVY)), P("Une ligne par véhicule, dans l'onglet « Parc »."), ("", None),
  H("Colonnes obligatoires (en-tête marqué *)"),
  P("• Immatriculation, Agence, Solution, Type de véhicule. L'immatriculation identifie le véhicule (pas de numéro de parc)."),
  P("• Immatriculation au format GH148HB : lettres et chiffres collés, sans tiret ni espace (les tirets éventuels sont retirés à l'import)."),
  P("• Solution et Type de véhicule : liste déroulante. Agence : liste issue de l'onglet « Listes » (complétez-la), autre valeur possible."), ("", None),
  H("Équipements : remplis automatiquement selon la solution"),
  P("• Dès que la solution est choisie, chaque équipement affiche « Oui » (prévu pour cette solution) ou « — » (ne s'applique pas)."),
  P("• Seul geste à faire : remplacer « Oui » par « Non » si l'équipement manque sur ce véhicule."),
  P("• Pour revenir à la valeur automatique d'une case, recopiez la case du dessus (la formule est la même sur toute la colonne)."),
  P("• La correspondance solution / équipements se règle dans l'onglet « Listes » (tableau « Équipements par solution »)."), ("", None),
  H("Contrôles automatiques"),
  P("• Case rouge pâle : champ obligatoire manquant sur une ligne commencée.  • Texte rouge gras : immatriculation en double."),
  P("• « Non » sur fond orangé : équipement déclaré absent."), ("", None),
  H("Ligne d'exemple"),
  P("• La ligne 3 de l'onglet « Parc » (en gris) montre le format attendu : remplacez-la ou supprimez-la. Elle est ignorée à l'import (commentaire « EXEMPLE »)."), ("", None),
  ("Une fois rempli : dans l'outil Matériel embarqué, onglet Réglages › Parc de véhicules, glissez ce fichier dans la zone de dépôt, vérifiez le rapport puis validez.", F(bold=True, size=10, color=GREEN))]
for i, (t, f) in enumerate(lignes, start=2):
    c = me.cell(row=i, column=2, value=t); c.alignment = Alignment(wrap_text=True, vertical='top')
    if f: c.font = f

# ---------- Listes ----------
li = wb.create_sheet("Listes")
def tete(cell, t, w=None):
    c = li[cell]; c.value = t; c.font = F(bold=True, color='FFFFFF', size=10); c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    if w: li.column_dimensions[c.column_letter].width = w
tete('A1', "Solutions", 16); tete('B1', "Types de véhicule", 24); tete('C1', "Agences (à compléter)", 24)
for r, v in enumerate(SOLUTIONS, 2): li.cell(row=r, column=1, value=v).font = F(size=10)
for r, v in enumerate(TYPES, 2): li.cell(row=r, column=2, value=v).font = F(size=10)
for r in range(2, 102): c = li.cell(row=r, column=3); c.border = box; c.font = F(size=10)
li['C1'].comment = Comment("Saisissez vos agences (une par ligne) : elles alimentent la liste déroulante de l'onglet Parc.", "Modèle")
li.column_dimensions['D'].width = 3
tete('E1', "Équipements par solution", 16)
for j, e in enumerate(EQUIP):
    tete(li.cell(row=1, column=C_EQ + j).coordinate, e, 12)
for i, s in enumerate(SOLUTIONS):
    r = 2 + i; c = li.cell(row=r, column=C_EQ - 1, value=s); c.font = F(bold=True, size=10, color=NAVY); c.border = box
    for j in range(len(EQUIP)):
        c = li.cell(row=r, column=C_EQ + j, value="Oui" if (j < CABINE or s == "CITI'ID") else "—")
        c.font = F(size=10, color='0000FF'); c.alignment = Alignment(horizontal='center'); c.border = box
li.row_dimensions[1].height = 42
li['E7'] = "Valeurs en bleu modifiables : « Oui » = équipement prévu pour la solution, « — » = ne s'applique pas. Source : guide Simpliciti CITI'ID et consignes de l'exploitation."
li['E7'].font = F(size=9, italic=True, color='5F6B7A'); li.merge_cells('E7:M8'); li['E7'].alignment = Alignment(wrap_text=True, vertical='top')
d = DataValidation(type='list', formula1='"Oui,—"', allow_blank=False); d.add('F2:M5'); li.add_data_validation(d)
li.freeze_panes = 'A2'

# ---------- Parc ----------
ws = wb.create_sheet("Parc", 1)
cols = [("Immatriculation *", 17), ("Agence *", 20), ("Solution *", 15), ("Type de véhicule *", 24), ("Marque / modèle", 24)] + [(e, 12) for e in EQUIP] + [("Commentaire", 40)]
F_EQ, L_EQ, COM = C_EQ, C_EQ + len(EQUIP) - 1, len(cols)
groupes = [(1, 5, "VÉHICULE", NAVY), (F_EQ, F_EQ + CABINE - 1, "ÉQUIPEMENTS CABINE - auto selon la solution", GREEN),
           (F_EQ + CABINE, L_EQ, "ÉQUIPEMENTS ARRIÈRE - auto selon la solution", '5B6680'), (COM, COM, "", NAVY)]
for a, b, t, color in groupes:
    ws.cell(row=1, column=a, value=t)
    if b > a: ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
    for c in range(a, b + 1):
        x = ws.cell(row=1, column=c); x.fill = PatternFill('solid', fgColor=color); x.font = F(bold=True, size=9, color='FFFFFF'); x.alignment = Alignment(horizontal='center', vertical='center')
HR, R0 = 2, 3; last = R0 + N
for i, (h, w) in enumerate(cols, 1):
    c = ws.cell(row=HR, column=i, value=h); c.font = F(bold=True, size=10, color=NAVY); c.fill = PatternFill('solid', fgColor='E3E8F0')
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = box; ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 20; ws.row_dimensions[HR].height = 32
from openpyxl.utils import get_column_letter as L
for r in range(R0, last + 1):
    for i in range(1, COM + 1):
        c = ws.cell(row=r, column=i); c.border = box; c.font = F(size=10)
        if F_EQ <= i <= L_EQ:
            col = c.column_letter
            c.value = f'=IF($C{r}="","",IFERROR(INDEX(Listes!{col}$2:{col}$5,MATCH($C{r},Listes!$E$2:$E$5,0)),""))'
            c.alignment = Alignment(horizontal='center')
exemple = {1: "GH214AB", 2: "Lyon Est", 3: "CITI'ID", 4: "BOM 26 T", 5: "Renault D Wide", F_EQ + 6: "Non", COM: "EXEMPLE - ligne à remplacer ou supprimer (ici : pas de BBV montée)"}
for i in range(1, COM + 1):
    c = ws.cell(row=R0, column=i)
    if i in exemple: c.value = exemple[i]
    c.font = F(size=10, italic=True, color='8792AB'); c.fill = PatternFill('solid', fgColor=GREY)
ws.freeze_panes = ws.cell(row=R0, column=2)
ws.auto_filter.ref = f"A{HR}:{L(COM)}{last}"
def liste(formule, plage, titre, msg, strict=True):
    d = DataValidation(type='list', formula1=formule, allow_blank=True, showErrorMessage=True, errorStyle='stop' if strict else 'warning',
                       errorTitle=titre, error=msg, showInputMessage=True, promptTitle=titre, prompt=msg)
    d.add(plage); ws.add_data_validation(d)
liste("=Listes!$A$2:$A$5", f"C{R0}:C{last}", "Solution", "Choisissez la solution : les équipements se remplissent seuls.")
liste("=Listes!$B$2:$B$14", f"D{R0}:D{last}", "Type de véhicule", "Choisissez le type dans la liste (pictos Paprec).")
liste("=Listes!$C$2:$C$101", f"B{R0}:B{last}", "Agence", "Agence de la liste « Listes » ; une autre valeur est acceptée.", strict=False)
liste('"Oui,Non"', f"{L(F_EQ)}{R0}:{L(L_EQ)}{last}", "Équipement", "Rempli selon la solution. Mettez « Non » si l'équipement manque sur ce véhicule.", strict=False)
rouge = PatternFill('solid', fgColor='FBE3E3', end_color='FBE3E3')
for col in "ABCD":
    ws.conditional_formatting.add(f"{col}{R0}:{col}{last}", FormulaRule(formula=[f'AND(COUNTA($A{R0}:$E{R0})>0,LEN(TRIM({col}{R0}))=0)'], fill=rouge))
ws.conditional_formatting.add(f"A{R0}:A{last}", FormulaRule(formula=[f'AND(A{R0}<>"",COUNTIF($A${R0}:$A${last},A{R0})>1)'], font=Font(name='Arial', bold=True, color='C62828')))
eq = f"{L(F_EQ)}{R0}:{L(L_EQ)}{last}"
ws.conditional_formatting.add(eq, FormulaRule(formula=[f'{L(F_EQ)}{R0}="Non"'], fill=PatternFill('solid', fgColor='FFEBDC', end_color='FFEBDC'), font=Font(name='Arial', bold=True, color='AE571E')))
ws.conditional_formatting.add(eq, FormulaRule(formula=[f'{L(F_EQ)}{R0}="—"'], font=Font(name='Arial', color='AAB4C3')))
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '1:2'
wb.active = 1
wb.save(SORTIE); print(SORTIE)
